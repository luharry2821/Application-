"""Excel-backed datastore for Ava (Skyline Car Rental).

A single workbook (./data/skyline.xlsx) with four worksheets — Reservations,
Incidents, Roadside, LostAndFound — created with headers on first run. Every
tool the voice agent can call reads/writes here.

Why one workbook with tabs (not live Google Sheets): zero external setup. Open
it in Excel anytime to inspect or edit. To move to live Google Sheets later,
swap this module's read/append/update calls for gspread against a service
account — the rest of the app doesn't change.

Concurrency: openpyxl does a full read-modify-write per call, so a module-level
lock serializes writers. FastAPI runs the (sync) tool handlers in a threadpool,
so the lock is what keeps two overlapping calls from clobbering the file. Don't
keep the workbook open in Excel during a live call — Excel locks/overwrites it.
"""

from __future__ import annotations

import os
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

DATA_DIR = Path(os.environ.get("SKYLINE_DATA_DIR", "data"))
WORKBOOK_PATH = DATA_DIR / "skyline.xlsx"

# Worksheet name -> ordered column headers.
SHEETS: dict[str, list[str]] = {
    "Reservations": [
        "confirmation_id", "created_at", "status", "phone", "drivers_license",
        "email", "car_type", "pickup_date", "pickup_time", "dropoff_date",
        "dropoff_time", "protection", "extras", "days", "daily_rate", "total",
        "notes",
    ],
    "Incidents": [
        "incident_id", "created_at", "phone", "drivers_license",
        "what_happened", "location", "occurred_at", "drivable", "status",
    ],
    "Roadside": [
        "case_id", "created_at", "phone", "drivers_license", "issue_type",
        "location", "eta_minutes", "status",
    ],
    "LostAndFound": [
        "case_id", "created_at", "phone", "drivers_license", "item_description",
        "location", "status",
    ],
}

_lock = threading.RLock()


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def new_id(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:6].upper()}"


def norm_phone(value) -> str:
    """Digits only, so '(555) 123-4567' and '555-123-4567' match."""
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def norm_license(value) -> str:
    return "".join(str(value or "").split()).upper()


def ensure_workbook() -> None:
    """Create the workbook (and any missing tabs/headers) if needed."""
    with _lock:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        if WORKBOOK_PATH.exists():
            wb = load_workbook(WORKBOOK_PATH)
            changed = False
            for name, headers in SHEETS.items():
                if name not in wb.sheetnames:
                    ws = wb.create_sheet(name)
                    ws.append(headers)
                    changed = True
            # openpyxl seeds a default "Sheet"; drop it if it's empty and unused.
            if "Sheet" in wb.sheetnames and "Sheet" not in SHEETS:
                default = wb["Sheet"]
                if default.max_row == 1 and default.max_column == 1 and default["A1"].value is None:
                    wb.remove(default)
                    changed = True
            if changed:
                wb.save(WORKBOOK_PATH)
            return

        wb = Workbook()
        wb.remove(wb.active)  # drop the auto-created blank sheet
        for name, headers in SHEETS.items():
            ws = wb.create_sheet(name)
            ws.append(headers)
        wb.save(WORKBOOK_PATH)


def _read_rows(ws, headers: list[str]) -> list[dict]:
    """Return data rows as dicts, tagging each with its 1-based _row index."""
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue
        record = {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers)}
        record["_row"] = idx
        rows.append(record)
    return rows


def append(sheet: str, values: dict) -> dict:
    """Append a row to `sheet`, filling only known columns. Returns the row written."""
    headers = SHEETS[sheet]
    with _lock:
        wb = load_workbook(WORKBOOK_PATH)
        ws = wb[sheet]
        ws.append([values.get(h, "") for h in headers])
        wb.save(WORKBOOK_PATH)
    return {h: values.get(h, "") for h in headers}


def find_reservation(phone: str, drivers_license: str) -> dict | None:
    """Most recent non-cancelled reservation matching phone + license, else None."""
    p, lic = norm_phone(phone), norm_license(drivers_license)
    with _lock:
        wb = load_workbook(WORKBOOK_PATH)
        rows = _read_rows(wb["Reservations"], SHEETS["Reservations"])
    matches = [
        r for r in rows
        if norm_phone(r.get("phone")) == p
        and norm_license(r.get("drivers_license")) == lic
        and str(r.get("status", "")).lower() != "cancelled"
    ]
    return matches[-1] if matches else None


def update_reservation(row_index: int, updates: dict, append_note: str | None = None) -> dict:
    """Patch named columns on a reservation row; optionally append to notes."""
    headers = SHEETS["Reservations"]
    with _lock:
        wb = load_workbook(WORKBOOK_PATH)
        ws = wb["Reservations"]
        for col, header in enumerate(headers, start=1):
            if header in updates and updates[header] not in (None, ""):
                ws.cell(row=row_index, column=col, value=updates[header])
        if append_note:
            note_col = headers.index("notes") + 1
            cell = ws.cell(row=row_index, column=note_col)
            existing = str(cell.value or "").strip()
            cell.value = f"{existing} | {append_note}".strip(" |") if existing else append_note
        wb.save(WORKBOOK_PATH)
        result = {h: ws.cell(row=row_index, column=i + 1).value for i, h in enumerate(headers)}
    return result
